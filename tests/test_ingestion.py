import openpyxl
import pytest
import yaml

from bridgework.ingestion import BridgeworkSchemaError, load_from_excel, load_from_yaml, write_excel
from bridgework.model import Drivers

V1 = Drivers(
    revenue_prior=4527.0,
    revenue_growth=0.031,
    cogs_pct=0.240,
    opex_pct=0.520,
    tax_rate=0.180,
    nwc_pct=0.050,
    capex_pct=0.030,
)


# ---------------------------------------------------------------- #
# YAML
# ---------------------------------------------------------------- #
def test_load_from_yaml_round_trip(tmp_path):
    p = tmp_path / "v1.yaml"
    p.write_text(
        yaml.dump({"schema_version": "1.0.0", "label": "test", "drivers": V1.to_dict()})
    )
    loaded = load_from_yaml(p)
    assert loaded == V1


def test_load_from_yaml_missing_drivers_key(tmp_path):
    p = tmp_path / "bad.yaml"
    p.write_text(yaml.dump({"schema_version": "1.0.0"}))
    with pytest.raises(BridgeworkSchemaError, match="drivers"):
        load_from_yaml(p)


def test_load_from_yaml_schema_major_mismatch(tmp_path):
    p = tmp_path / "bad_version.yaml"
    p.write_text(yaml.dump({"schema_version": "2.0.0", "drivers": V1.to_dict()}))
    with pytest.raises(BridgeworkSchemaError, match="major mismatch"):
        load_from_yaml(p)


def test_load_from_yaml_missing_driver_role(tmp_path):
    d = V1.to_dict()
    del d["capex_pct"]
    p = tmp_path / "incomplete.yaml"
    p.write_text(yaml.dump({"schema_version": "1.0.0", "drivers": d}))
    with pytest.raises(BridgeworkSchemaError):
        load_from_yaml(p)


# ---------------------------------------------------------------- #
# Excel round-trip
# ---------------------------------------------------------------- #
def test_write_then_load_excel_round_trip_fidelity(tmp_path):
    p = tmp_path / "v1.xlsx"
    write_excel(V1, p, "V1 Test")
    loaded = load_from_excel(p)
    for role in V1.to_dict():
        assert getattr(loaded, role) == pytest.approx(getattr(V1, role), abs=1e-9)


def test_excel_round_trip_permits_driver_value_edit(tmp_path):
    p = tmp_path / "v1.xlsx"
    write_excel(V1, p, "V1 Test")
    wb = openpyxl.load_workbook(p)
    wb["Drivers"]["B5"] = 0.10  # revenue_growth
    wb.save(p)
    loaded = load_from_excel(p)
    assert loaded.revenue_growth == pytest.approx(0.10, abs=1e-9)


def test_excel_round_trip_rejects_calc_tampering(tmp_path):
    p = tmp_path / "v1.xlsx"
    write_excel(V1, p, "V1 Test")
    wb = openpyxl.load_workbook(p)
    wb["Calc"]["B7"] = "=B5*0.99"  # tamper with Opex formula
    wb.save(p)
    with pytest.raises(BridgeworkSchemaError, match="modified outside the permitted round-trip scope"):
        load_from_excel(p)


def test_excel_round_trip_rejects_missing_sheet(tmp_path):
    p = tmp_path / "v1.xlsx"
    write_excel(V1, p, "V1 Test")
    wb = openpyxl.load_workbook(p)
    del wb["Meta"]
    wb.save(p)
    with pytest.raises(BridgeworkSchemaError, match="missing required sheet"):
        load_from_excel(p)


def test_excel_round_trip_rejects_schema_version_major_mismatch(tmp_path):
    p = tmp_path / "v1.xlsx"
    write_excel(V1, p, "V1 Test")
    wb = openpyxl.load_workbook(p)
    wb["Meta"]["B1"] = "9.0.0"
    wb.save(p)
    with pytest.raises(BridgeworkSchemaError, match="major mismatch"):
        load_from_excel(p)


def test_excel_round_trip_rejects_row_order_shuffle(tmp_path):
    p = tmp_path / "v1.xlsx"
    write_excel(V1, p, "V1 Test")
    wb = openpyxl.load_workbook(p)
    ws = wb["Drivers"]
    ws["A4"], ws["A5"] = ws["A5"].value, ws["A4"].value
    wb.save(p)
    with pytest.raises(BridgeworkSchemaError, match="row order mismatch"):
        load_from_excel(p)


def test_excel_generated_workbook_has_no_hardcoded_calc_cells(tmp_path):
    """Every Calc cell must be a live formula, never a Python-computed
    literal (Blueprint §B LOCKED: Excel is a byproduct, not the source of
    truth)."""
    p = tmp_path / "v1.xlsx"
    write_excel(V1, p, "V1 Test")
    wb = openpyxl.load_workbook(p, data_only=False)
    calc = wb["Calc"]
    for row in range(4, 14):
        val = calc.cell(row=row, column=2).value
        assert isinstance(val, str) and val.startswith("="), f"row {row} is not a formula: {val!r}"
