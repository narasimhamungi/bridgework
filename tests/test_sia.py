import openpyxl
import pytest

from bridgework.ingestion import write_excel
from bridgework.model import Drivers
from bridgework.sia import run_structural_checks

V1 = Drivers(
    revenue_prior=4527.0,
    revenue_growth=0.031,
    cogs_pct=0.240,
    opex_pct=0.520,
    tax_rate=0.180,
    nwc_pct=0.050,
    capex_pct=0.030,
)


@pytest.fixture
def clean_workbook(tmp_path):
    p = tmp_path / "clean.xlsx"
    write_excel(V1, p, "Clean")
    return p


def test_all_checks_pass_on_clean_workbook(clean_workbook):
    report = run_structural_checks(clean_workbook)
    assert report.passed
    assert len(report.results) == 8


def test_check1_schema_version_mismatch_fails(clean_workbook):
    wb = openpyxl.load_workbook(clean_workbook)
    wb["Meta"]["B1"] = "9.9.9"
    wb.save(clean_workbook)
    report = run_structural_checks(clean_workbook)
    result = next(r for r in report.results if r.check_id == 1)
    assert not result.passed
    assert not report.passed


def test_check2_missing_sheet_fails(clean_workbook):
    wb = openpyxl.load_workbook(clean_workbook)
    del wb["Meta"]
    wb.save(clean_workbook)
    report = run_structural_checks(clean_workbook)
    result = next(r for r in report.results if r.check_id == 2)
    assert not result.passed


def test_check3_hardcoded_literal_fails(clean_workbook):
    wb = openpyxl.load_workbook(clean_workbook)
    wb["Calc"]["B7"] = 1234.5  # hardcoded number instead of formula
    wb.save(clean_workbook)
    report = run_structural_checks(clean_workbook)
    result = next(r for r in report.results if r.check_id == 3)
    assert not result.passed
    assert not report.passed


def test_check4_broken_cross_reference_fails(clean_workbook):
    wb = openpyxl.load_workbook(clean_workbook)
    wb["Calc"]["B7"] = "=B5*Drivers!$B$99"  # out-of-range reference
    wb.save(clean_workbook)
    report = run_structural_checks(clean_workbook)
    result = next(r for r in report.results if r.check_id == 4)
    assert not result.passed


def test_check5_sign_convention_fails(clean_workbook):
    wb = openpyxl.load_workbook(clean_workbook)
    wb["Drivers"]["B4"] = -100.0  # revenue_prior must be > 0
    wb.save(clean_workbook)
    report = run_structural_checks(clean_workbook)
    result = next(r for r in report.results if r.check_id == 5)
    assert not result.passed


def test_check6_row_order_fails(clean_workbook):
    wb = openpyxl.load_workbook(clean_workbook)
    ws = wb["Drivers"]
    ws["A4"], ws["A5"] = ws["A5"].value, ws["A4"].value
    wb.save(clean_workbook)
    report = run_structural_checks(clean_workbook)
    result = next(r for r in report.results if r.check_id == 6)
    assert not result.passed


def test_check7_unsupported_construct_fails(clean_workbook):
    wb = openpyxl.load_workbook(clean_workbook)
    wb["Calc"]["B7"] = "=INDIRECT(\"B5\")*Drivers!$B$7"
    wb.save(clean_workbook)
    report = run_structural_checks(clean_workbook)
    result = next(r for r in report.results if r.check_id == 7)
    assert not result.passed


def test_check8_orphaned_driver_detected(clean_workbook):
    """A driver declared but never referenced in Calc is dead input a
    reader would wrongly assume is live. The Blueprint promised this check
    and the code did not implement it until adversarial review found the
    drift."""
    wb = openpyxl.load_workbook(clean_workbook)
    # Point the Opex formula at a literal so opex_pct becomes unreferenced.
    wb["Calc"]["B7"] = "=B5*0.5"
    wb.save(clean_workbook)
    report = run_structural_checks(clean_workbook)
    result = next(r for r in report.results if r.check_id == 8)
    assert not result.passed
    assert "opex_pct" in result.message


def test_check8_passes_on_clean_workbook(clean_workbook):
    report = run_structural_checks(clean_workbook)
    assert next(r for r in report.results if r.check_id == 8).passed


def test_sia_is_a_regression_gate_not_a_tamper_boundary(clean_workbook):
    """Documents the security boundary precisely, because the Blueprint
    previously implied SIA catches hand-edits and it does not.

    A semantically tampered formula using only whitelisted constructs and
    resolvable references passes SIA, and is caught by ingestion's
    round-trip diff. Locking this in prevents the claim drifting back."""
    from bridgework.ingestion import BridgeworkSchemaError, load_from_excel

    wb = openpyxl.load_workbook(clean_workbook)
    wb["Calc"]["B10"] = "=B8-B10*2"  # valid syntax, wrong economics
    wb.save(clean_workbook)

    assert run_structural_checks(clean_workbook).passed, (
        "SIA is expected to PASS this — it validates syntax, not semantics"
    )
    with pytest.raises(BridgeworkSchemaError, match="modified outside the permitted round-trip scope"):
        load_from_excel(clean_workbook)


def test_sia_report_summary_is_readable(clean_workbook):
    report = run_structural_checks(clean_workbook)
    summary = report.summary()
    assert "SIA: PASS" in summary
    assert "8/8" in summary
